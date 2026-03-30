#!/usr/bin/env python3
"""Generate an offer-only PartialUpdate file for existing Amazon listings."""
import csv
import glob
import os
import re
from datetime import datetime

import openpyxl

# Pricing config (same as ali_to_amazon.py)
DEFAULT_PRICE_GBP = 12.99
TARGET_PROFIT_MARGIN = 0.30
AMAZON_REFERRAL_FEE = 0.1545
AMAZON_PER_ITEM_FEE = 0.75
USD_TO_GBP = 0.79
ALI_SHIPPING_ESTIMATE = 2.00
MIN_SELL_PRICE = 5.99


def parse_price(price_str):
    if not price_str or price_str == "N/A":
        return None
    cleaned = re.sub(r'[^\d.,]', '', price_str)
    if ',' in cleaned and '.' not in cleaned:
        cleaned = cleaned.replace(',', '.')
    elif ',' in cleaned and '.' in cleaned:
        cleaned = cleaned.replace(',', '')
    try:
        return float(cleaned)
    except ValueError:
        return None


def ali_to_gbp(price_usd):
    if not price_usd:
        return DEFAULT_PRICE_GBP
    cost_gbp = price_usd * USD_TO_GBP
    total_cost = cost_gbp + ALI_SHIPPING_ESTIMATE
    denominator = 1 - AMAZON_REFERRAL_FEE - TARGET_PROFIT_MARGIN
    if denominator <= 0:
        return DEFAULT_PRICE_GBP
    sell_price = (total_cost + AMAZON_PER_ITEM_FEE) / denominator
    sell_price = round(sell_price, 2)
    actual_profit = sell_price - (sell_price * AMAZON_REFERRAL_FEE) - AMAZON_PER_ITEM_FEE - total_cost
    if actual_profit < 6.0:
        sell_price = (6.0 + AMAZON_PER_ITEM_FEE + total_cost) / (1 - AMAZON_REFERRAL_FEE)
        sell_price = round(sell_price, 2)
    if sell_price < MIN_SELL_PRICE:
        sell_price = MIN_SELL_PRICE
    return sell_price


def main():
    print("Loading template...")
    wb = openpyxl.load_workbook('TOY_FIGURE_TOYS_AND_GAMES.xlsm', keep_vba=True)
    ws = wb['Template']

    # Build column map
    col_map = {}
    for c in range(1, 460):
        val = ws.cell(row=3, column=c).value
        if val:
            col_map[str(val).strip().lower()] = c

    # Columns needed for offer update
    needed_fields = [
        'feed_product_type',
        'item_sku',
        'update_delete',
        'condition_type',
        'fulfillment_availability#1.fulfillment_channel_code',
        'fulfillment_availability#1.quantity',
        'fulfillment_availability#1.lead_time_to_ship_max_days',
        'purchasable_offer[marketplace_id=a1f83g8c2aro7p]#1.our_price#1.schedule#1.value_with_tax',
        # Required listing fields — Amazon won't activate offers on
        # listings missing these, even if the offer data is accepted.
        'country_of_origin',
        'batteries_required',
        'are_batteries_included',
        'supplier_declared_dg_hz_regulation1',
    ]

    offer_cols = []
    for field in needed_fields:
        c = col_map.get(field.lower())
        if c:
            offer_cols.append((field, c))
            print(f"  {field} -> col {c}")
        else:
            print(f"  MISSING: {field}")

    # Get headers
    row1_vals = [str(ws.cell(row=1, column=c).value or "") for _, c in offer_cols]
    row2_vals = [str(ws.cell(row=2, column=c).value or "") for _, c in offer_cols]
    row3_vals = [str(ws.cell(row=3, column=c).value or "") for _, c in offer_cols]

    wb.close()

    # Load prices from CSVs
    sku_prices = {}
    csvs = sorted(glob.glob('*resin_models.csv'), key=os.path.getmtime, reverse=True)
    for csv_path in csvs:
        try:
            with open(csv_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    pid = row.get("id", "").strip()
                    sku = "ALI-" + pid
                    price_str = row.get("product_price", "")
                    if sku not in sku_prices:
                        sku_prices[sku] = price_str
        except Exception as e:
            print(f"  Error reading {csv_path}: {e}")
    print(f"Loaded prices for {len(sku_prices)} SKUs")

    # ALL 50 SKUs from the listing
    all_skus = [
        "ALI-1005008646642221", "ALI-1005008409205488", "ALI-1005004703422009",
        "ALI-1005008355199678", "ALI-1005006916914032", "ALI-1005008481590562",
        "ALI-1005007046881571", "ALI-1005008662078125", "ALI-1005009864244373",
        "ALI-1005008599403323", "ALI-1005007172577952", "ALI-1005010244881345",
        "ALI-1005009136087654", "ALI-1005009361332814", "ALI-1005008380114065",
        "ALI-1005005695747888", "ALI-1005010285882352", "ALI-1005010487348414",
        "ALI-1005008515354917", "ALI-1005008310733727", "ALI-1005009592570857",
        "ALI-1005008952753669", "ALI-1005008311225625", "ALI-1005010569629065",
        "ALI-1005008608454206", "ALI-1005008481588304", "ALI-1005009657571555",
        "ALI-1005009281795534", "ALI-1005009452566364", "ALI-1005007136214211",
        "ALI-1005010084413384", "ALI-1005008500713411", "ALI-1005007888200911",
        "ALI-1005010750673658", "ALI-1005008542918421", "ALI-1005004276931602",
        "ALI-1005008252638147", "ALI-1005009584000185", "ALI-1005008515696288",
        "ALI-1005008971488155", "ALI-1005008414221558", "ALI-1005008247046418",
        "ALI-1005009281636927", "ALI-1005003372896446", "ALI-1005008306140627",
        "ALI-1005009361601101", "ALI-1005007446770017", "ALI-1005008281706783",
        "ALI-1005008708230536", "ALI-1005008608478387",
    ]

    # Generate file
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_name = f"amazon_offer_update_{ts}.txt"

    with open(output_name, "w", encoding="utf-8") as f:
        f.write("\t".join(row1_vals) + "\n")
        f.write("\t".join(row2_vals) + "\n")
        f.write("\t".join(row3_vals) + "\n")

        for sku in all_skus:
            price_str = sku_prices.get(sku, "")
            price_usd = parse_price(price_str)
            sell_price = ali_to_gbp(price_usd)

            row_data = []
            for field, c in offer_cols:
                if field == 'feed_product_type':
                    row_data.append('toyfigure')
                elif field == 'item_sku':
                    row_data.append(sku)
                elif field == 'update_delete':
                    row_data.append('PartialUpdate')
                elif field == 'condition_type':
                    row_data.append('New')
                elif 'fulfillment_channel_code' in field:
                    row_data.append('DEFAULT')
                elif 'quantity' in field:
                    row_data.append('5')
                elif 'lead_time' in field:
                    row_data.append('7')
                elif 'our_price' in field:
                    row_data.append(str(sell_price))
                elif field == 'country_of_origin':
                    row_data.append('China')
                elif field == 'batteries_required':
                    row_data.append('No')
                elif field == 'are_batteries_included':
                    row_data.append('No')
                elif 'supplier_declared_dg_hz_regulation' in field:
                    row_data.append('Not Applicable')
                else:
                    row_data.append('')
            f.write("\t".join(row_data) + "\n")

    print(f"\nDone! Generated: {output_name}")
    print(f"  {len(all_skus)} SKUs with PartialUpdate + offer data")
    print(f"\nUpload this file via: Catalogue > Add Products via Upload")
    print(f"This will add offer/price data to your existing listings.")


if __name__ == "__main__":
    main()
